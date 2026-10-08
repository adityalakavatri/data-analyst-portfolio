"""
Builds excel/Ecommerce_Sales_Summary.xlsx from the cleaned SQL output.

Sheets
  Sales_Data   - month x category x region table (formatted as an Excel Table, ready for a PivotTable)
  Summary      - SUMIFS / margin formulas by category, region and year (live formulas)
  Pivot_Guide  - step-by-step instructions to build the PivotTable + slicers

Run after analysis.py:  python python/build_excel.py
"""
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
con = sqlite3.connect(ROOT / "data" / "ecommerce.db")
df = pd.read_sql_query(
    """
    SELECT order_month AS Month, CAST(SUBSTR(order_month, 1, 4) AS INT) AS Year,
           category AS Category, region AS Region,
           COUNT(DISTINCT order_id) AS Orders,
           ROUND(SUM(revenue), 0) AS Revenue, ROUND(SUM(profit), 0) AS Profit
    FROM sales_fact WHERE status = 'Delivered'
    GROUP BY order_month, category, region
    ORDER BY order_month, category, region
    """,
    con,
)

F = "Arial"
bold, head_fill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
base = Font(name=F, size=10)

wb = Workbook()

# ------------------------------------------------------------------ Sales_Data
ws = wb.active
ws.title = "Sales_Data"
ws.append(list(df.columns))
for row in df.itertuples(index=False):
    ws.append(list(row))
for c in ws[1]:
    c.font, c.fill = bold, head_fill
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.font = base
    r[5].number_format = r[6].number_format = '#,##0'
n = len(df) + 1
tab = Table(displayName="SalesTbl", ref=f"A1:G{n}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tab)
for col, w in zip("ABCDEFG", [11, 8, 18, 10, 9, 14, 14]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
rng = lambda col: f"Sales_Data!${col}$2:${col}${n}"

# ------------------------------------------------------------------ Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "E-commerce sales summary (Delivered orders, INR)"
sm["A1"].font = Font(name=F, bold=True, size=14)
sm["A2"] = "All numbers below are live SUMIFS formulas on the Sales_Data sheet."
sm["A2"].font = Font(name=F, italic=True, size=9, color="666666")


def block(top, title, key_col, keys):
    """Writes a title row, header row, one SUMIFS row per key, and a total row."""
    sm.cell(top, 1, title).font = Font(name=F, bold=True, size=11)
    for j, h in enumerate([title.split(" by ")[1].title(), "Revenue", "Profit", "Margin %", "Share of revenue"], 1):
        c = sm.cell(top + 1, j, h)
        c.font, c.fill = bold, head_fill
    first = top + 2
    last = first + len(keys) - 1
    for i, k in enumerate(keys):
        r = first + i
        sm.cell(r, 1, k)
        sm.cell(r, 2, f"=SUMIFS({rng('F')},{rng(key_col)},A{r})")
        sm.cell(r, 3, f"=SUMIFS({rng('G')},{rng(key_col)},A{r})")
        sm.cell(r, 4, f"=IF(B{r}=0,0,C{r}/B{r})")
        sm.cell(r, 5, f"=IF(B${last + 1}=0,0,B{r}/B${last + 1})")
    t = last + 1
    sm.cell(t, 1, "Total")
    sm.cell(t, 2, f"=SUM(B{first}:B{last})")
    sm.cell(t, 3, f"=SUM(C{first}:C{last})")
    sm.cell(t, 4, f"=IF(B{t}=0,0,C{t}/B{t})")
    sm.cell(t, 5, f"=SUM(E{first}:E{last})")
    for r in range(first, t + 1):
        for j in range(1, 6):
            c = sm.cell(r, j)
            c.font = Font(name=F, size=10, bold=(r == t))
        sm.cell(r, 2).number_format = sm.cell(r, 3).number_format = '#,##0'
        sm.cell(r, 4).number_format = sm.cell(r, 5).number_format = '0.0%'
    return first, last


cats = sorted(df["Category"].unique())
regs = sorted(df["Region"].unique())
years = [int(y) for y in sorted(df["Year"].unique())]
c1, c2 = block(4, "Revenue by category", "C", cats)
block(c2 + 4, "Revenue by region", "D", regs)
block(c2 + 4 + len(regs) + 4, "Revenue by year", "B", years)

ch = BarChart()
ch.type, ch.title, ch.style = "col", "Revenue by category (INR)", 10
ch.add_data(Reference(sm, min_col=2, min_row=5, max_row=c2), titles_from_data=True)
ch.set_categories(Reference(sm, min_col=1, min_row=c1, max_row=c2))
ch.legend, ch.height, ch.width = None, 7.5, 14
sm.add_chart(ch, "H4")
for col, w in zip("ABCDE", [24, 16, 16, 11, 17]):
    sm.column_dimensions[col].width = w

# ------------------------------------------------------------------ Pivot_Guide
pg = wb.create_sheet("Pivot_Guide")
steps = [
    "HOW TO BUILD THE PIVOT TABLE (takes ~2 minutes) - this is the Excel skill to show in interviews",
    "",
    "1. Open the Sales_Data sheet and click any cell inside the table (it is already an Excel Table named SalesTbl).",
    "2. Insert > PivotTable > New Worksheet.",
    "3. Rows: Category. Columns: Year. Values: Sum of Revenue and Sum of Profit.",
    "4. Add a calculated field: Margin = Profit / Revenue (PivotTable Analyze > Fields, Items & Sets > Calculated Field).",
    "5. Insert > Slicer > Region and Month so stakeholders can filter interactively.",
    "6. Sort the Category rows by Sum of Revenue (largest first) and apply Conditional Formatting > Data Bars.",
    "7. Insert > PivotChart (clustered column) next to the table and label the axes in INR.",
    "",
    "Interview talking point: the Summary sheet cross-checks the PivotTable - both must return identical totals.",
]
for i, s in enumerate(steps, 1):
    c = pg.cell(i, 1, s)
    c.font = Font(name=F, bold=(i == 1), size=10)
pg.column_dimensions["A"].width = 120

out = ROOT / "excel" / "Ecommerce_Sales_Summary.xlsx"
wb.save(out)
print("saved", out)
