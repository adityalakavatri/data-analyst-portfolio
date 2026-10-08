"""
Builds excel/HR_Attrition_Dashboard.xlsx from the cleaned employee data.

Sheets
  Employee_Data - cleaned data as an Excel Table (ready for PivotTables)
  Summary       - COUNTIFS-based attrition rates by department, overtime, satisfaction, tenure (live formulas)
  Pivot_Guide   - how to build the PivotTable + slicers

Run after analysis.py:  python python/build_excel.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
df = pd.read_csv(ROOT / "data" / "processed" / "employees_clean.csv")
# helper column so tenure bands can be counted with plain COUNTIFS
df["tenure_band"] = pd.cut(df["years_at_company"], [-1, 2, 5, 10, 100],
                           labels=["0-2 yrs", "3-5 yrs", "6-10 yrs", "10+ yrs"]).astype(str)

F = "Arial"
hdr_font, hdr_fill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
body = Font(name=F, size=10)

wb = Workbook()
ws = wb.active
ws.title = "Employee_Data"
ws.append(list(df.columns))
for row in df.itertuples(index=False):
    ws.append(list(row))
for c in ws[1]:
    c.font, c.fill = hdr_font, hdr_fill
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.font = body
n = len(df) + 1
last_col = ws.cell(1, len(df.columns)).column_letter
t = Table(displayName="EmpTbl", ref=f"A1:{last_col}{n}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(t)
ws.freeze_panes = "A2"
for i, col in enumerate(df.columns, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(col) + 2)

colletter = {c: ws.cell(1, i).column_letter for i, c in enumerate(df.columns, 1)}
R = lambda c: f"Employee_Data!${colletter[c]}$2:${colletter[c]}${n}"

sm = wb.create_sheet("Summary")
sm["A1"] = "HR attrition summary (live COUNTIFS formulas on Employee_Data)"
sm["A1"].font = Font(name=F, bold=True, size=14)

# overall KPIs
sm["A3"], sm["B3"] = "Employees", f"=COUNTA({R('employee_id')})"
sm["A4"], sm["B4"] = "Leavers", f'=COUNTIF({R("attrition")},"Yes")'
sm["A5"], sm["B5"] = "Attrition rate", "=IF(B3=0,0,B4/B3)"
sm["A6"], sm["B6"] = "Avg monthly income (INR)", f"=AVERAGE({R('monthly_income')})"
sm["B5"].number_format = "0.0%"
sm["B6"].number_format = "#,##0"
for r in range(3, 7):
    sm.cell(r, 1).font = Font(name=F, bold=True, size=10)
    sm.cell(r, 2).font = body


def block(top, title, col, keys):
    sm.cell(top, 1, title).font = Font(name=F, bold=True, size=11)
    for j, h in enumerate([col.replace("_", " ").title(), "Employees", "Leavers", "Attrition %"], 1):
        c = sm.cell(top + 1, j, h)
        c.font, c.fill = hdr_font, hdr_fill
    for i, k in enumerate(keys):
        r = top + 2 + i
        sm.cell(r, 1, k)
        sm.cell(r, 2, f"=COUNTIFS({R(col)},A{r})")
        sm.cell(r, 3, f'=COUNTIFS({R(col)},A{r},{R("attrition")},"Yes")')
        sm.cell(r, 4, f"=IF(B{r}=0,0,C{r}/B{r})")
        sm.cell(r, 4).number_format = "0.0%"
        for j in range(1, 5):
            sm.cell(r, j).font = body
    return top + 2, top + 1 + len(keys)


deps = sorted(df["department"].unique())
d1, d2 = block(8, "Attrition by department", "department", deps)
o1, o2 = block(d2 + 3, "Attrition by overtime", "overtime", ["No", "Yes"])
s1, s2 = block(o2 + 3, "Attrition by job satisfaction (1=low, 4=high)", "job_satisfaction", [1, 2, 3, 4])
t1, t2 = block(s2 + 3, "Attrition by tenure band", "tenure_band", ["0-2 yrs", "3-5 yrs", "6-10 yrs", "10+ yrs"])

ch = BarChart()
ch.type, ch.title, ch.style = "col", "Attrition % by department", 10
ch.add_data(Reference(sm, min_col=4, min_row=d1 - 1, max_row=d2), titles_from_data=True)
ch.set_categories(Reference(sm, min_col=1, min_row=d1, max_row=d2))
ch.legend, ch.height, ch.width = None, 7.5, 13
sm.add_chart(ch, "G3")
for col, w in zip("ABCD", [34, 12, 12, 12]):
    sm.column_dimensions[col].width = w

pg = wb.create_sheet("Pivot_Guide")
lines = [
    "HOW TO BUILD THE PIVOT TABLE (about 2 minutes)",
    "",
    "1. On Employee_Data click any cell inside the table (named EmpTbl). Insert > PivotTable > New Worksheet.",
    "2. Rows: department. Values: Count of employee_id.",
    "3. Drag attrition to Columns, then right-click a value > Show Values As > % of Row Total.",
    "4. Insert slicers for overtime, job_satisfaction and tenure_band. Click them to see attrition change instantly.",
    "5. Conditional Formatting > Color Scales on the 'Yes' column to highlight hot-spots.",
    "6. Insert > PivotChart (stacked 100% column).",
    "",
    "Cross-check: the PivotTable totals must match the Summary sheet and the SQL output (data/processed/*.csv).",
]
for i, s in enumerate(lines, 1):
    pg.cell(i, 1, s).font = Font(name=F, bold=(i == 1), size=10)
pg.column_dimensions["A"].width = 115

out = ROOT / "excel" / "HR_Attrition_Dashboard.xlsx"
wb.save(out)
print("saved", out)
